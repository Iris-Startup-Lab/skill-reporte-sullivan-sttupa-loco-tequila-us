"""
================================================================================
 DERIVA account_mapping.csv DESDE EL LIBRO DE REFERENCIA DEL CLIENTE
================================================================================
Herramienta de un solo uso (se vuelve a correr cuando el cliente manda un corte
nuevo). Lee las hojas de cuentas del libro `Sales Report Loco USA <fecha>.xlsx` y
escribe `Config/loco_tequila_us/account_mapping.csv` con cuenta -> vendedor ->
canal.

POR QUÉ EXISTE
--------------
La bandera `--account-map` está cableada en el orquestador, en el procesador y en
el generador de PDF desde septiembre, y **nunca se le suministró un archivo**. Ese
es el motivo real de que el leaderboard de vendedores estuviera dominado por una
barra "Unknown": no era un bug de cálculo, era un insumo ausente. El mapa que hacía
falta ya existía, dentro del propio libro del cliente.

PRECEDENCIA ENTRE HOJAS
-----------------------
`Accounts H2-<año>` gana sobre `Accounts H1-<año>`, y ambas sobre
`Account Summary Total Business`. El orden no es arbitrario: el cliente coloca H2
antes de H1 en su libro precisamente porque es el corte más reciente, y
`Account Summary` es más amplia pero acumula histórico.

DÓNDE VIVE LA SALIDA, Y POR QUÉ NO EN Client_Data/
--------------------------------------------------
`Config/` es dato de referencia **derivado y editable**. `Client_Data/` es materia
prima confidencial y además los lectores la recorren con glob, así que un CSV ahí
corre riesgo de ser recogido como si fuera un insumo y se confunde con dato
transaccional.

USO
---
    python Scripts/tools/derive_loco_account_map.py --workbook "<ruta al .xlsx>"
    python Scripts/tools/derive_loco_account_map.py --workbook "..." --dry-run
================================================================================
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

import openpyxl

_HERE = Path(__file__).resolve()
PROJECT_ROOT = _HERE.parent.parent.parent
if str(PROJECT_ROOT / "Scripts") not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT / "Scripts"))

from loco_attribution import normalize_account  # noqa: E402

OUT_PATH = PROJECT_ROOT / "Config" / "loco_tequila_us" / "account_mapping.csv"

# (hoja, col cuenta, col vendedor, col canal, primera fila de datos)
# Los índices son 0-based sobre la fila cruda. Se validan por encabezado antes de
# leer: si el cliente reordena columnas, se falla en voz alta en vez de escribir
# un mapa silenciosamente equivocado.
SHEET_SPECS = [
    ("Accounts H2-",  0, 8, 9, 2, ("Salesperson", "Channel")),
    ("Accounts H1-",  0, 8, 9, 2, ("Salesperson", "Channel")),
    ("Account Summary Total Business", 0, 12, 11, 2, ("Salesperson", "Channel")),
]

# Vocabulario de canal del cliente -> el nuestro. "Off and On" es su valor crudo
# para las cuentas que operan de las dos formas.
CHANNEL_MAP = {
    "on premise": "On Premise",
    "off premise": "Off Premise",
    "off and on": "Combo",
    "combo": "Combo",
    "dtc": "DTC",
    "unknown": "Unknown",
}


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in {"nan", "nat", "none", "-", ""} else s


def _validate_header(row, cs, cc, expected, sheet) -> None:
    got = (_clean(row[cs]) if cs < len(row) else "",
           _clean(row[cc]) if cc < len(row) else "")
    if got != expected:
        raise SystemExit(
            f"[ERROR] La hoja '{sheet}' no tiene las columnas esperadas en su "
            f"posición: se esperaba {expected} y se encontró {got}. El cliente "
            "reordenó o renombró columnas; ajusta SHEET_SPECS antes de seguir, "
            "porque escribir el mapa a ciegas produciría atribuciones falsas."
        )


def derive(workbook: Path) -> list[dict]:
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    seen: dict[str, dict] = {}
    for prefix, ca, cs, cc, start, expected in SHEET_SPECS:
        name = next((s for s in wb.sheetnames if s.startswith(prefix)), None)
        if not name:
            print(f"  [AVISO] No se encontró una hoja que empiece con '{prefix}'.")
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        if len(rows) <= start:
            continue
        _validate_header(rows[start - 1], cs, cc, expected, name)
        added = 0
        for r in rows[start:]:
            account = _clean(r[ca]) if ca < len(r) else ""
            if not account:
                continue
            key = normalize_account(account)
            if not key or key in seen:
                continue
            sp = _clean(r[cs]) if cs < len(r) else ""
            ch = _clean(r[cc]) if cc < len(r) else ""
            ch_norm = CHANNEL_MAP.get(ch.casefold(), ch or "Unknown")
            # Una fila sin vendedor NI canal utilizable no aporta nada al mapa.
            if not sp and ch_norm in ("", "Unknown"):
                continue
            seen[key] = {
                "account": account,
                "salesperson": sp or "Unknown",
                "channel": ch_norm or "Unknown",
                "source_sheet": name,
            }
            added += 1
        print(f"  {name:34s} +{added:4d} cuentas")
    wb.close()
    return sorted(seen.values(), key=lambda r: r["account"].casefold())


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--workbook", required=True,
                    help="Libro de referencia del cliente (.xlsx).")
    ap.add_argument("--out", default=str(OUT_PATH))
    ap.add_argument("--dry-run", action="store_true",
                    help="Solo reporta; no escribe el CSV.")
    args = ap.parse_args()

    wb_path = Path(args.workbook)
    if not wb_path.exists():
        raise SystemExit(f"[ERROR] No existe el libro: {wb_path}")

    print(f"Derivando mapa de cuentas desde: {wb_path.name}")
    rows = derive(wb_path)
    if not rows:
        raise SystemExit("[ERROR] No se derivó ninguna cuenta. Revisa SHEET_SPECS.")

    from collections import Counter
    print(f"\n  total: {len(rows)} cuentas")
    print(f"  vendedores: {dict(Counter(r['salesperson'] for r in rows).most_common())}")
    print(f"  canales:    {dict(Counter(r['channel'] for r in rows).most_common())}")

    if args.dry_run:
        print("\n  [dry-run] No se escribió nada.")
        return 0

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["account", "salesperson", "channel",
                                           "source_sheet"])
        w.writeheader()
        w.writerows(rows)
    print(f"\n  Escrito: {out}")
    return 0


if __name__ == "__main__":
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError, OSError):
            pass
    raise SystemExit(main())
